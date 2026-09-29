"""The cover letter, as the PDF the submission system asks for.

Kept separate from the manuscript build because the letter is not part of the manuscript and
must not acquire its figures, headings or numbering. Same Word export settings, for the same
reason: background repagination off.

    python paper/build_cover_letter.py
"""

from __future__ import annotations

import subprocess
from pathlib import Path

import docx
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
SOURCE = HERE / "cover_letter.md"
BUILD = HERE / "build"
DOCX = BUILD / "cover_letter.docx"
PDF = BUILD / "PMB_cover_letter.pdf"


def main() -> int:
    text = SOURCE.read_text(encoding="utf-8")
    document = docx.Document()
    for section in document.sections:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)
        for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, margin, Inches(1.0))

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    normal.paragraph_format.space_after = Pt(10)

    for block in text.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        # The signature block is several short lines that must not be rejoined.
        if "\n" in block and max(len(line) for line in block.splitlines()) < 70:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            for position, line in enumerate(block.splitlines()):
                if position:
                    paragraph.add_run().add_break()
                paragraph.add_run(line.strip())
            continue
        document.add_paragraph(" ".join(block.split()))

    BUILD.mkdir(parents=True, exist_ok=True)
    document.save(DOCX)

    script = BUILD / "export_letter.ps1"
    script.write_text(
        "$ErrorActionPreference = 'Stop'\n"
        "$word = New-Object -ComObject Word.Application\n"
        "$word.Visible = $false\n"
        "$word.DisplayAlerts = 0\n"
        "$word.Options.Pagination = $false\n"
        f"$document = $word.Documents.Open('{DOCX}', $false, $true)\n"
        f"$document.ExportAsFixedFormat('{PDF}', 17)\n"
        "$document.Close(0)\n"
        "$word.Quit()\n",
        encoding="utf-8")
    result = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script)],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=600)
    if result.returncode or not PDF.is_file():
        print(result.stdout[-1200:])
        print(result.stderr[-1200:])
        raise SystemExit("cover letter PDF export failed")

    print(f"  {PDF.relative_to(HERE.parent)}  ({PDF.stat().st_size // 1024} kB)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
