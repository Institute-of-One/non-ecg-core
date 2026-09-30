"""Turn the manuscript into the single PDF the journal wants, with every number resolved.

The manuscript is the source; the .docx and the .pdf are products and are never edited by
hand. Markers are resolved here from paper/frozen/manifest.json, so the document cannot carry
a number that no result file produced.

Deliberately not included: page numbers and line numbers. The submission system stamps its
own, and a page-number field in the footer is what made Word's PDF export fail to terminate
on a figure-heavy manuscript in this programme. `Options.Pagination` is switched off in the
export for the same reason.

    python paper/build_submission.py
"""

from __future__ import annotations

import functools
import json
import re
import subprocess
import sys
from pathlib import Path

import docx

# The console here is cp932 and cannot print replacement characters.
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor

HERE = Path(__file__).resolve().parent
MANUSCRIPT = HERE / "manuscript.md"
MANIFEST = HERE / "frozen" / "manifest.json"
BUILD = HERE / "build"
DOCX = BUILD / "manuscript_pmb.docx"
PDF = BUILD / "PMB_manuscript.pdf"
RESOLVED_MD = BUILD / "manuscript_resolved.md"

FIGURE_WIDTH = Inches(6.1)
BODY_FONT = "Times New Roman"
BODY_SIZE = Pt(11)


def resolve_markers(text: str, manifest: dict) -> str:
    unresolved: list[str] = []

    def replace(match: re.Match) -> str:
        path = match.group(1).split(":")[-1].split(".")
        try:
            return str(functools.reduce(lambda node, key: node[key], path, manifest))
        except (KeyError, TypeError):
            unresolved.append(match.group(1))
            return match.group(0)

    out = re.sub(r"\[\[results:([^\]]+)\]\]", replace, text)
    if unresolved:
        raise SystemExit(f"unresolved markers: {sorted(set(unresolved))}")
    return out


#: Bold first and non-greedily, so that a bold span containing italics — **the value of
#: *n*min is a judgement** — is recognised instead of falling through as literal asterisks.
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<!\*)\*[^*\n]+\*(?!\*))")


def add_runs(paragraph, text: str, bold: bool = False, italic: bool = False) -> None:
    """Inline **bold**, *italic* and `code`, with italics allowed inside bold."""
    for piece in INLINE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            add_runs(paragraph, piece[2:-2], bold=True, italic=italic)
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9.5)
            run.bold, run.italic = bold, italic
        elif piece.startswith("*") and piece.endswith("*") and len(piece) > 2:
            run = paragraph.add_run(piece[1:-1])
            run.bold, run.italic = bold, True
        else:
            run = paragraph.add_run(piece)
            run.bold, run.italic = bold, italic


def build_docx(markdown: str, destination=None, expected_figures: int = 6) -> None:
    document = docx.Document()

    for section in document.sections:
        section.page_width, section.page_height = Inches(8.27), Inches(11.69)   # A4
        for margin in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
            setattr(section, margin, Inches(1.0))

    normal = document.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)

    lines = markdown.splitlines()
    index = 0
    figures = 0
    while index < len(lines):
        line = lines[index].rstrip()
        index += 1

        if not line.strip():
            continue

        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
        if image:
            path = (HERE / image.group(2)).resolve()
            if not path.is_file():
                raise SystemExit(f"figure missing: {path}")
            document.add_picture(str(path), width=FIGURE_WIDTH)
            document.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            # Figure 5's caption broke across two pages, so the reader met the plot on one
            # page and the legend on the next.
            document.paragraphs[-1].paragraph_format.keep_with_next = True
            figures += 1
            continue

        if line.startswith("### ") or line.startswith("## "):
            level = 2 if line.startswith("### ") else 1
            heading = document.add_heading("", level=level)
            # Headings carry emphasis too, and passing the raw text printed *N*min literally.
            add_runs(heading, line.split(" ", 1)[1].strip())
            for run in heading.runs:                     # Word's blue is not manuscript style
                run.font.color.rgb = RGBColor(0, 0, 0)
            continue
        if line.startswith("# "):
            heading = document.add_heading(line[2:].strip(), level=0)
            for run in heading.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
            continue
        if line.strip() == "---":
            continue

        # A paragraph continues until a blank line, so wrapped source lines rejoin.
        block = [line]
        while index < len(lines) and lines[index].strip() and not re.match(
                r"^(#{1,3} |!\[|> |\d+\. |---$)", lines[index]):
            block.append(lines[index].rstrip())
            index += 1
        text = " ".join(block).strip()

        if text.startswith("> "):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            add_runs(paragraph, text[2:])
            for run in paragraph.runs:
                run.italic = True
            continue

        numbered = re.match(r"^(\d+)\. (.*)$", text)
        if numbered:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.35)
            paragraph.paragraph_format.first_line_indent = Inches(-0.35)
            paragraph.paragraph_format.space_after = Pt(4)
            add_runs(paragraph, f"{numbered.group(1)}. {numbered.group(2)}")
            continue

        paragraph = document.add_paragraph()
        # A caption set in the body face at the body size reads as body text, and the
        # reader has to work out where the figure's description stops.
        if text.startswith("**Figure "):
            fmt = paragraph.paragraph_format
            fmt.left_indent = Inches(0.3)
            fmt.right_indent = Inches(0.3)
            fmt.space_before = Pt(4)
            fmt.space_after = Pt(14)
            fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
            fmt.keep_together = True             # and not split across a page break itself
            add_runs(paragraph, text)
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
            continue
        add_runs(paragraph, text)

    if figures != expected_figures:
        raise SystemExit(
            f"expected {expected_figures} figures in the document, placed {figures}")

    BUILD.mkdir(parents=True, exist_ok=True)
    target = destination or DOCX
    document.save(target)
    print(f"  {target.relative_to(HERE.parent)}  ({figures} figures)")


def export_pdf() -> None:
    """Word, with background repagination off. See the module docstring."""
    # Export beside the target and move it into place. Word refuses to overwrite a PDF that
    # a viewer has open, and reports it as a COM exception that says nothing useful; this
    # way the failure is named where it happens.
    staged = BUILD / "PMB_manuscript.staged.pdf"
    staged.unlink(missing_ok=True)
    script = BUILD / "export_pdf.ps1"
    script.write_text(
        "$ErrorActionPreference = 'Stop'\n"
        "$word = New-Object -ComObject Word.Application\n"
        "try {\n"
        "  $word.Visible = $false\n"
        "  $word.DisplayAlerts = 0\n"
        "  $word.Options.Pagination = $false\n"
        "  $word.Options.CheckSpellingAsYouType = $false\n"
        "  $word.Options.CheckGrammarAsYouType = $false\n"
        f"  $document = $word.Documents.Open('{DOCX}', $false, $true)\n"
        f"  $document.ExportAsFixedFormat('{staged}', 17)\n"
        "  $document.Close(0)\n"
        "} finally {\n"
        # An export that throws must still close Word, or the next build finds its own
        # .docx locked by the instance the last failure left behind.
        "  $word.Quit()\n"
        "}\n",
        encoding="utf-8")
    result = subprocess.run(
        ["powershell", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File", str(script)],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=900)
    if result.returncode or not staged.is_file():
        print((result.stdout or "")[-1500:])
        print((result.stderr or "")[-1500:])
        raise SystemExit("PDF export failed")
    try:
        staged.replace(PDF)
    except PermissionError:
        raise SystemExit(
            f"{PDF.name} is open in another application, so the new build could not "
            f"replace it. Close it and run this again; the new file is at {staged}"
        ) from None
    print(f"  {PDF.relative_to(HERE.parent)}  ({PDF.stat().st_size // 1024} kB)")


def main() -> int:
    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    text = MANUSCRIPT.read_text(encoding="utf-8")
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    text = resolve_markers(text, manifest)

    BUILD.mkdir(parents=True, exist_ok=True)
    RESOLVED_MD.write_text(text, encoding="utf-8")
    print(f"  {RESOLVED_MD.relative_to(HERE.parent)}")

    build_docx(text)
    export_pdf()
    print("\nbuilt. Open the PDF and look at it before it goes anywhere.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
