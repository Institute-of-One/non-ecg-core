"""Does the built PDF satisfy what this journal asks for, and does it match its source?

A submission kit copied from the last paper asserted three things that were false about the
next one. So nothing here is asserted: every rule is read out of the built PDF or out of the
manuscript, and the journal's limits are named as constants that can be checked against the
guide.

Physics in Medicine and Biology, Paper: at most 10,000 words and 12 figures, a structured
abstract under the four headings below, and references with article titles. The submission
system stamps its own page and line numbers, so the built document carries neither.

    python paper/presubmission_check.py
"""

from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

import docx
import fitz

# The console here is cp932 and cannot print an em dash.
sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
PDF = HERE / "build" / "PMB_manuscript.pdf"
RESOLVED = HERE / "build" / "manuscript_resolved.md"
MANUSCRIPT = HERE / "manuscript.md"
MANIFEST = HERE / "frozen" / "manifest.json"
RESOLVED_REFS = HERE / "frozen" / "references_resolved.json"

MAX_WORDS = 10000
MAX_FIGURES = 12
EXPECTED_FIGURES = 6
MAX_ABSTRACT_WORDS = 300
ABSTRACT_HEADINGS = ("Objective", "Approach", "Main results", "Significance")
REQUIRED_DECLARATIONS = ("Funding", "Competing interests", "Ethical statement",
                         "Data and code availability", "Use of generative AI",
                         "Prior contact")


def words(text: str) -> int:
    return len(re.findall(r"[A-Za-z0-9'–‐-]+", text))


def main() -> int:
    failures: list[str] = []

    def check(label: str, passed: bool, detail: str = "") -> None:
        print(f"  {'PASS' if passed else 'FAIL'}  {label}{'  | ' + detail if detail else ''}")
        if not passed:
            failures.append(label)

    for path in (PDF, RESOLVED):
        if not path.is_file():
            print(f"  FAIL  {path.name} is missing; run paper/build_submission.py first")
            return 1

    source = MANUSCRIPT.read_text(encoding="utf-8")
    resolved = RESOLVED.read_text(encoding="utf-8")
    document = fitz.open(PDF)
    pdf_text = "\n".join(page.get_text() for page in document)

    check("no unresolved result marker in the built document",
          "[[results:" not in resolved and "[[results:" not in pdf_text)
    check("no editorial placeholder left",
          "[DECIDE:" not in resolved and "[CITE:" not in resolved,
          "DECIDE" if "[DECIDE:" in resolved else "")

    body = resolved.split("## References")[0]
    check(f"body within {MAX_WORDS} words", words(body) <= MAX_WORDS, str(words(body)))

    abstract = resolved.split("## Abstract")[1].split("---")[0]
    check(f"abstract within {MAX_ABSTRACT_WORDS} words",
          words(abstract) <= MAX_ABSTRACT_WORDS, str(words(abstract)))
    for heading in ABSTRACT_HEADINGS:
        check(f"abstract carries the heading {heading!r}", f"**{heading}.**" in abstract)

    images = sum(len(page.get_images(full=True)) for page in document)
    captions = len(re.findall(r"\*\*Figure \d+\.\*\*", resolved))
    check(f"at most {MAX_FIGURES} figures", images <= MAX_FIGURES, f"{images} embedded")
    check("every figure is embedded in the PDF", images == EXPECTED_FIGURES, str(images))
    check("every figure has a caption", captions == EXPECTED_FIGURES, str(captions))

    mentions = [int(n) for n in re.findall(r"\(figure (\d+)\)", body)]
    check("figures first mentioned in numerical order",
          mentions == sorted(set(mentions)) == list(range(1, EXPECTED_FIGURES + 1)),
          str(mentions))

    # The journal stamps its own furniture, and a page-number field in a footer is what made
    # Word's export fail to terminate here before. Checked on the .docx, where a footer
    # either holds text or does not.
    built = docx.Document(HERE / "build" / "manuscript_pmb.docx")
    furniture = [part.text.strip() for section in built.sections
                 for part in (section.header, section.footer) for part in part.paragraphs]
    check("no header or footer text in the built document", not any(furniture),
          str([f for f in furniture if f]))

    # Bold sentences in running text are the author's standing objection, and they kept coming
    # back because nothing checked for them: the existing check only confirmed that the "**"
    # markers had been consumed by the builder, not that the emphasis was wanted. A journal uses
    # bold for the structured-abstract headings, the caption label, the declaration labels and
    # Harvard volume numbers. Everywhere else it is the author's voice, not the journal's.
    allowed = re.compile(r"^(Objective|Approach|Main results|Significance)\.$|"
                         r"^Figure S?\d+\.$|"
                         r"^(Author and affiliation|Funding|Competing interests|"
                         r"Ethical statement|Data and code availability|Use of generative AI|"
                         r"Prior contact)\.$|^\d+$")
    stray = []
    in_references = False
    for paragraph in built.paragraphs:
        if paragraph.text.strip().startswith("References"):
            in_references = True
        if in_references:
            continue
        for run in paragraph.runs:
            text = run.text.strip()
            if run.bold and text and not allowed.match(text):
                stray.append(f"{paragraph.text[:40]}… -> {text[:60]}")
    check("no bold emphasis in running text", not stray, str(stray[:4]))

    references = json.loads(RESOLVED_REFS.read_text(encoding="utf-8"))
    listing = resolved.split("## References")[1]
    check("every resolved reference appears in the list",
          all(r["doi"].lower() in listing.lower() for r in references),
          f"{len(references)} references")
    check("every reference carries an article title",
          all(len(r["title"]) > 10 for r in references))

    for declaration in REQUIRED_DECLARATIONS:
        check(f"declaration present: {declaration}", f"**{declaration}" in resolved)

    # The source wraps at 90 columns, so the affiliation can straddle a line break.
    flat = re.sub(r"\s+", " ", resolved)
    check("affiliation is the canonical string",
          "Institute of One, LISIT Co., Ltd., Tokyo 150-0044, Japan" in flat)
    check("ORCID present", "0000-0001-9211-1071" in resolved)

    # check_references.py refuses to write the reference section when a citation does not
    # match, and leaves the previous one in place. A rewrapped citation therefore failed there
    # while every check here passed against a reference list belonging to the last good build.
    source_citations = set(re.findall(r"\(([A-Z][A-Za-zÀ-ɏ'-]+(?: (?:et al|and "
                                      r"[A-Z][A-Za-zÀ-ɏ'-]+))?) (\d{4})\)", source))
    flat_source = re.sub(r"\s+", " ", source)
    wrapped = {f"{who} {year}" for who, year in
               re.findall(r"\(([A-Z][A-Za-zÀ-ɏ'-]+(?: (?:et al|and "
                          r"[A-Z][A-Za-zÀ-ɏ'-]+))?)\s+(\d{4})\)", flat_source)
               if (who, year) not in source_citations}
    check("no citation split across a line in the source", not wrapped, str(sorted(wrapped)))

    # The data-availability statement names a release and promises that a clean copy of it
    # rebuilds every number here. What has to hold is not that the tag is HEAD — the manuscript
    # goes on being edited, and the DOI it cites cannot be inside the snapshot it names — but
    # that nothing which produces a number has moved since the tag was cut. v0.1.0 was named by
    # a manuscript four rounds of revision newer than it, and that is what this catches.
    PRODUCES_NUMBERS = ("analysis", "results", "paper/frozen", "paper/collect_results.py",
                        "paper/figures.py")
    named = sorted(set(re.findall(r"\bv\d+\.\d+\.\d+\b", source)))
    for tag in named:
        commit = subprocess.run(["git", "rev-list", "-n1", tag], cwd=REPO,
                                capture_output=True, text=True).stdout.strip()
        if not commit:
            check(f"release {tag} exists and still produces this manuscript's numbers",
                  False, "no such tag")
            continue
        ancestor = subprocess.run(["git", "merge-base", "--is-ancestor", tag, "HEAD"],
                                  cwd=REPO, capture_output=True).returncode == 0
        moved = subprocess.run(["git", "diff", "--name-only", tag, "--", *PRODUCES_NUMBERS],
                               cwd=REPO, capture_output=True, text=True).stdout.split()
        check(f"release {tag} exists and still produces this manuscript's numbers",
              ancestor and not moved,
              "the tag is not an ancestor of HEAD" if not ancestor else str(sorted(moved)))
    check("the manuscript names the release it is submitted with", bool(named), str(named))

    manifest = json.loads(MANIFEST.read_text(encoding="utf-8"))
    used = {m.split(":")[-1].split(".")[-1]
            for m in re.findall(r"\[\[results:([^\]]+)\]\]", source)}
    check("every marker the source uses exists in the manifest",
          used <= set(manifest["metrics"]), str(sorted(used - set(manifest["metrics"]))))

    print()
    print(f"pages: {document.page_count} | body words: {words(body)} | "
          f"abstract words: {words(abstract)} | figures: {images} | "
          f"references: {len(references)}")
    if failures:
        print(f"\nNOT READY: {len(failures)} check(s) failed")
        return 1
    print("\nALL GREEN — and now open the PDF and look at it")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
